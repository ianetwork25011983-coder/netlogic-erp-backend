"""
Módulo 14 - Reportes (Excel).

Cada función devuelve un `openpyxl.Workbook` ya armado; las vistas son
las que lo serializan a bytes y lo devuelven como descarga. Mantener la
generación separada de la vista permite reusar estas funciones desde
Celery (ej. un reporte mensual programado) sin pasar por HTTP.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.billing.models import DocumentoVenta
from apps.inventory.services import InventoryService

HEADER_FILL = PatternFill(start_color="1F2A37", end_color="1F2A37", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _set_header(ws, headers):
    ws.append(headers)
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"


def _autosize(ws, headers):
    for col_idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(14, len(str(header)) + 4)


class ReportService:
    @staticmethod
    def kardex_excel(producto, deposito=None, lote=None):
        wb = Workbook()
        ws = wb.active
        ws.title = "Kardex"
        headers = [
            "Fecha", "Tipo de Movimiento", "Depósito", "Cantidad", "Costo Unitario PYG",
            "Costo Total PYG", "Saldo Cantidad", "Saldo Valor PYG", "Documento", "Usuario",
        ]
        _set_header(ws, headers)

        movimientos = InventoryService.get_kardex(producto, deposito=deposito, lote=lote)
        for m in movimientos:
            ws.append([
                m.fecha.replace(tzinfo=None), m.get_tipo_movimiento_display(), str(m.deposito),
                float(m.cantidad), float(m.costo_unitario_pyg), float(m.costo_total_pyg),
                float(m.saldo_cantidad_posterior), float(m.saldo_valor_posterior_pyg),
                f"{m.documento_tipo} {m.documento_referencia}".strip(), str(m.usuario),
            ])
        _autosize(ws, headers)
        return wb

    @staticmethod
    def inventario_valorizado_excel(empresa, deposito=None, categoria=None):
        wb = Workbook()
        ws = wb.active
        ws.title = "Inventario Valorizado"
        headers = ["Código", "Producto", "Depósito", "Cantidad", "Costo Promedio PYG", "Valor Total PYG"]
        _set_header(ws, headers)

        saldos = InventoryService.get_valorizacion(empresa=empresa, deposito=deposito, categoria=categoria)
        total_general = 0
        for s in saldos.select_related("producto", "deposito"):
            ws.append([
                s.producto.codigo, s.producto.nombre, str(s.deposito),
                float(s.cantidad), float(s.costo_promedio_pyg), float(s.valor_total_pyg),
            ])
            total_general += float(s.valor_total_pyg)

        ws.append([])
        ws.append(["", "", "", "", "TOTAL GENERAL", total_general])
        _autosize(ws, headers)
        return wb

    @staticmethod
    def ventas_excel(empresa, fecha_desde, fecha_hasta):
        wb = Workbook()
        ws = wb.active
        ws.title = "Ventas"
        headers = [
            "Fecha", "Tipo", "Número", "Cliente", "Moneda", "Subtotal",
            "Impuesto", "Descuento Global", "Total", "Total PYG", "Estado",
        ]
        _set_header(ws, headers)

        documentos = DocumentoVenta.objects.filter(
            empresa=empresa, fecha_emision__gte=fecha_desde, fecha_emision__lte=fecha_hasta,
        ).select_related("cliente", "moneda").order_by("fecha_emision")

        for d in documentos:
            ws.append([
                d.fecha_emision, d.get_tipo_documento_display(), d.numero,
                d.cliente.nombre_comercial or d.cliente.razon_social, d.moneda.code,
                float(d.subtotal), float(d.impuesto_total), float(d.descuento_global),
                float(d.total), float(d.total_pyg), d.get_estado_display(),
            ])
        _autosize(ws, headers)
        return wb

    @staticmethod
    def compras_excel(empresa, fecha_desde, fecha_hasta):
        from apps.purchases.models import OrdenCompra

        wb = Workbook()
        ws = wb.active
        ws.title = "Compras"
        headers = ["Fecha", "Número OC", "Proveedor", "Moneda", "Total", "Estado"]
        _set_header(ws, headers)

        ordenes = OrdenCompra.objects.filter(
            empresa=empresa, fecha__date__gte=fecha_desde, fecha__date__lte=fecha_hasta,
        ).select_related("proveedor", "moneda").order_by("fecha")

        for o in ordenes:
            ws.append([
                o.fecha.date(), o.numero, o.proveedor.nombre_comercial or o.proveedor.razon_social,
                o.moneda.code, float(o.total), o.get_estado_display(),
            ])
        _autosize(ws, headers)
        return wb

    @staticmethod
    def clientes_excel(empresa):
        from apps.customers.models import Cliente

        wb = Workbook()
        ws = wb.active
        ws.title = "Clientes"
        headers = ["RUC", "Razón Social", "Nombre Comercial", "Teléfono", "Email", "Días Crédito", "Límite Crédito", "Activo"]
        _set_header(ws, headers)

        for c in Cliente.objects.filter(empresa=empresa).order_by("razon_social"):
            ws.append([
                c.ruc, c.razon_social, c.nombre_comercial, c.telefono, c.email,
                c.dias_credito, float(c.limite_credito), "Sí" if c.active else "No",
            ])
        _autosize(ws, headers)
        return wb

    @staticmethod
    def proveedores_excel(empresa):
        from apps.suppliers.models import Proveedor

        wb = Workbook()
        ws = wb.active
        ws.title = "Proveedores"
        headers = ["RUC", "Razón Social", "Nombre Comercial", "Teléfono", "Email", "Moneda Default", "Activo"]
        _set_header(ws, headers)

        for p in Proveedor.objects.filter(empresa=empresa).select_related("moneda_default").order_by("razon_social"):
            ws.append([
                p.ruc, p.razon_social, p.nombre_comercial, p.telefono, p.email,
                p.moneda_default.code, "Sí" if p.active else "No",
            ])
        _autosize(ws, headers)
        return wb

    @staticmethod
    def rentabilidad_excel(empresa, fecha_desde, fecha_hasta):
        from apps.billing.models import DocumentoVentaItem

        wb = Workbook()
        ws = wb.active
        ws.title = "Rentabilidad por Producto"
        headers = ["Código", "Producto", "Cantidad Vendida", "Ventas Netas PYG", "Costo PYG", "Margen Bruto PYG", "Margen %"]
        _set_header(ws, headers)

        items = DocumentoVentaItem.objects.filter(
            documento__empresa=empresa, documento__tipo_documento=DocumentoVenta.TIPO_FACTURA,
            documento__estado=DocumentoVenta.ESTADO_EMITIDA,
            documento__fecha_emision__gte=fecha_desde, documento__fecha_emision__lte=fecha_hasta,
        ).select_related("producto", "documento", "movimiento_inventario")

        acumulado = {}
        for item in items:
            key = item.producto_id
            registro = acumulado.setdefault(key, {
                "codigo": item.producto.codigo, "nombre": item.producto.nombre,
                "cantidad": 0, "ventas_pyg": 0.0, "costo_pyg": 0.0,
            })
            registro["cantidad"] += float(item.cantidad)
            registro["ventas_pyg"] += float(item.subtotal_linea * item.documento.tipo_cambio_pyg)
            if item.movimiento_inventario_id:
                registro["costo_pyg"] += float(item.movimiento_inventario.costo_total_pyg)

        for registro in sorted(acumulado.values(), key=lambda r: r["ventas_pyg"], reverse=True):
            margen = registro["ventas_pyg"] - registro["costo_pyg"]
            margen_pct = (margen / registro["ventas_pyg"] * 100) if registro["ventas_pyg"] else 0
            ws.append([
                registro["codigo"], registro["nombre"], registro["cantidad"],
                registro["ventas_pyg"], registro["costo_pyg"], margen, round(margen_pct, 2),
            ])
        _autosize(ws, headers)
        return wb
